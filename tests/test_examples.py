import csv
import importlib
import subprocess
import sys
from pathlib import Path

from openpyxl import load_workbook


def test_generated_examples_are_empty_and_do_not_expose_business_asins(tmp_path):
    generator = importlib.import_module("scripts.generate_examples")

    paths = generator.generate_examples(tmp_path)

    assert {path.name for path in paths} == {
        "asin_input_template.xlsx",
        "candidates_example.csv",
        "example_roi_output.xlsx",
    }

    input_workbook = load_workbook(tmp_path / "asin_input_template.xlsx", read_only=True)
    try:
        assert input_workbook.properties.creator == "asin-1688-roi-tool"
        asin_rows = list(input_workbook["ASIN输入"].iter_rows(values_only=True))
        assert len(asin_rows) == 1
    finally:
        input_workbook.close()

    with (tmp_path / "candidates_example.csv").open(
        "r", encoding="utf-8-sig", newline=""
    ) as handle:
        candidate_rows = list(csv.reader(handle))
    assert len(candidate_rows) == 1

    output_workbook = load_workbook(tmp_path / "example_roi_output.xlsx", read_only=True)
    try:
        assert output_workbook["ROI结果"].max_row == 1
        assert output_workbook["1688候选"].max_row == 1
        assert output_workbook["ASIN输入"].max_row == 1
    finally:
        output_workbook.close()


def test_example_generator_runs_from_source_checkout(tmp_path):
    project_root = Path(__file__).resolve().parents[1]

    completed = subprocess.run(
        [
            sys.executable,
            "-m",
            "scripts.generate_examples",
            "--output-dir",
            str(tmp_path),
        ],
        cwd=project_root,
        capture_output=True,
        text=True,
        check=False,
    )

    assert completed.returncode == 0, completed.stderr
    assert (tmp_path / "asin_input_template.xlsx").exists()
