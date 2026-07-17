from openpyxl import load_workbook

from asin_1688_roi.exporter import write_workbook
from asin_1688_roi.models import CandidateProduct, TargetProduct
from asin_1688_roi.roi import calculate_results
from asin_1688_roi.roi_config import RoiAssumptions


def test_exported_untrusted_text_is_not_written_as_excel_formula(tmp_path):
    output = tmp_path / "result.xlsx"
    candidate = CandidateProduct(
        asin="B000000001",
        search_keyword="测试",
        title="=1+1",
    )

    write_workbook(output, targets=[], candidates=[candidate], results=[])

    workbook = load_workbook(output, read_only=True, data_only=False)
    try:
        assert workbook.properties.creator == "asin-1688-roi-tool"
        cell = workbook["1688候选"]["D2"]
        assert cell.value == "=1+1"
        assert cell.data_type == "s"
    finally:
        workbook.close()


def test_export_records_actual_assumptions_source_and_snapshot(tmp_path):
    output = tmp_path / "result.xlsx"
    target = TargetProduct(asin="B000000001", selling_price_usd=20, cpc_usd=1)
    assumptions = RoiAssumptions(commission_rate=0.12)
    results = calculate_results(
        [target],
        [],
        7,
        assumptions=assumptions,
        assumptions_source="roi_profile.json",
    )

    write_workbook(output, targets=[target], candidates=[], results=results)

    workbook = load_workbook(output, read_only=True, data_only=False)
    try:
        roi_values = list(workbook["ROI结果"].values)
        assert "参数来源" in roi_values[0]
        assert roi_values[1][roi_values[0].index("参数来源")] == "roi_profile.json"

        parameter_rows = dict(workbook["参数说明"].values)
        assert parameter_rows["本次参数来源"] == "roi_profile.json"
        assert '"commission_rate":0.12' in parameter_rows["本次参数快照"]
    finally:
        workbook.close()
