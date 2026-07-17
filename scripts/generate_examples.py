from __future__ import annotations

import argparse
import csv
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

PROJECT_ROOT = Path(__file__).resolve().parents[1]
SOURCE_ROOT = PROJECT_ROOT / "src"
if str(SOURCE_ROOT) not in sys.path:
    sys.path.insert(0, str(SOURCE_ROOT))


def generate_examples(output_dir: str | Path) -> list[Path]:
    from asin_1688_roi.exporter import CANDIDATE_HEADERS, TARGET_HEADERS, write_workbook

    directory = Path(output_dir)
    directory.mkdir(parents=True, exist_ok=True)

    input_path = directory / "asin_input_template.xlsx"
    workbook = Workbook()
    workbook.properties.creator = "asin-1688-roi-tool"
    workbook.properties.lastModifiedBy = "asin-1688-roi-tool"
    workbook.properties.title = "ASIN input template"
    input_sheet = workbook.active
    input_sheet.title = "ASIN输入"
    input_sheet.append(TARGET_HEADERS)
    input_sheet.freeze_panes = "A2"
    for cell in input_sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F4E78")
    input_sheet.auto_filter.ref = f"A1:{input_sheet.cell(1, len(TARGET_HEADERS)).coordinate}"

    instructions = workbook.create_sheet("填写说明")
    instructions.append(["字段", "填写要求"])
    instructions.append(["ASIN", "10位字母或数字；模板不附带任何真实业务 ASIN。"])
    instructions.append(["1688搜索关键词", "至少填写1个中文关键词；多个关键词使用 | 分隔。"])
    instructions.append(["材质/功能/结构", "用于同规格匹配，不要只填写泛化品类名。"])
    instructions.append(["完整销售单元采购价", "在候选表人工确认，不能直接使用搜索页最低价。"])
    instructions.append(["汇率", "正式复核建议使用命令行 --fx 固定并留痕。"])
    instructions.column_dimensions["A"].width = 24
    instructions.column_dimensions["B"].width = 72
    workbook.save(input_path)
    workbook.close()

    candidate_path = directory / "candidates_example.csv"
    with candidate_path.open("w", encoding="utf-8-sig", newline="") as handle:
        csv.writer(handle).writerow(CANDIDATE_HEADERS)

    output_path = directory / "example_roi_output.xlsx"
    write_workbook(output_path, targets=[], candidates=[], results=[])
    return [input_path, candidate_path, output_path]


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description="生成不含真实业务数据的空白示例文件")
    parser.add_argument("--output-dir", type=Path, default=PROJECT_ROOT / "examples")
    args = parser.parse_args(argv)
    generate_examples(args.output_dir)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
