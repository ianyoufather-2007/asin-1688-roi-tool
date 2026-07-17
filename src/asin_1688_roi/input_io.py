from __future__ import annotations

import csv
import re
from pathlib import Path

from openpyxl import load_workbook

from asin_1688_roi.models import CandidateProduct, TargetProduct
from asin_1688_roi.utils import clean_text, first_nonempty, split_keywords, to_float

TARGET_ALIASES = {
    "asin": ("ASIN", "asin", "竞品ASIN"),
    "selling_price_usd": ("售价", "售价USD", "Selling Price", "selling_price_usd"),
    "cpc_usd": ("CPC", "PPC", "cpc_usd"),
    "amazon_title": ("Amazon商品名称", "商品名称", "标题", "amazon_title"),
    "search_keywords": ("1688搜索关键词", "搜索关键词", "search_keywords"),
    "material": ("材质", "material"),
    "function": ("功能", "function"),
    "structure": ("结构", "structure"),
    "sales_unit": ("销售单元", "sales_unit"),
    "pack_quantity": ("套装数量", "pack_quantity"),
    "accessory_quantity": ("配件数量", "accessory_quantity"),
    "length_cm": ("产品长cm", "长度cm", "length_cm"),
    "width_cm": ("产品宽cm", "宽度cm", "width_cm"),
    "height_cm": ("产品高cm", "高度cm", "height_cm"),
    "product_weight_kg": ("产品重量kg", "product_weight_kg"),
    "package_length_cm": ("包装长cm", "package_length_cm"),
    "package_width_cm": ("包装宽cm", "package_width_cm"),
    "package_height_cm": ("包装高cm", "package_height_cm"),
    "package_weight_kg": ("包装重量kg", "package_weight_kg"),
    "fba_fee_usd": ("FBA", "FBA_USD", "fba_fee_usd"),
    "notes": ("备注", "notes"),
}

CANDIDATE_ALIASES = {
    "asin": ("ASIN", "asin"),
    "search_keyword": ("1688搜索关键词", "搜索关键词", "search_keyword"),
    "offer_id": ("Offer ID", "offer_id"),
    "title": ("1688商品标题", "商品标题", "title"),
    "product_url": ("1688商品链接", "商品链接", "product_url"),
    "image_url": ("商品图片", "image_url"),
    "supplier": ("供应商", "supplier"),
    "shop_url": ("店铺链接", "shop_url"),
    "province": ("省份", "province"),
    "city": ("城市", "city"),
    "page_price_text": ("页面价格", "page_price_text"),
    "comparison_price_cny": ("完整销售单元采购价", "采购价", "comparison_price_cny"),
    "min_order_quantity": ("起订量", "min_order_quantity"),
    "sales_text": ("销量", "sales_text"),
    "return_rate": ("回头率", "return_rate"),
    "material": ("材质", "material"),
    "function": ("功能", "function"),
    "structure": ("结构", "structure"),
    "sales_unit": ("销售单元", "sales_unit"),
    "pack_quantity": ("套装数量", "pack_quantity"),
    "accessory_quantity": ("配件数量", "accessory_quantity"),
    "length_cm": ("长cm", "length_cm"),
    "width_cm": ("宽cm", "width_cm"),
    "height_cm": ("高cm", "height_cm"),
    "detail_text": ("详情文本", "detail_text"),
    "screenshot_path": ("截图路径", "screenshot_path"),
    "source_timestamp": ("抓取时间", "source_timestamp"),
    "manual_valid": ("人工有效", "manual_valid"),
    "manual_reason": ("人工判断原因", "manual_reason"),
}

ASIN_PATTERN = re.compile(r"^[A-Z0-9]{10}$")


def _rows_from_xlsx(path: Path, preferred_sheet: str | None = None) -> list[dict[str, object]]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        if preferred_sheet and preferred_sheet in workbook.sheetnames:
            sheet = workbook[preferred_sheet]
        else:
            sheet = workbook[workbook.sheetnames[0]]
        values = list(sheet.iter_rows(values_only=True))
    finally:
        workbook.close()
    if not values:
        return []
    headers = [clean_text(value) for value in values[0]]
    return [
        dict(zip(headers, row, strict=False))
        for row in values[1:]
        if any(v is not None for v in row)
    ]


def _rows_from_csv(path: Path) -> list[dict[str, object]]:
    with path.open("r", encoding="utf-8-sig", newline="") as handle:
        delimiter = "\t" if path.suffix.casefold() == ".tsv" else ","
        return list(csv.DictReader(handle, delimiter=delimiter))


def _validate_asin(value: object) -> str:
    asin = clean_text(value).upper()
    if asin and not ASIN_PATTERN.fullmatch(asin):
        raise ValueError(f"ASIN 格式无效：{asin}")
    return asin


def _read_rows(path: Path, preferred_sheet: str | None = None) -> list[dict[str, object]]:
    if path.suffix.casefold() == ".xlsx":
        return _rows_from_xlsx(path, preferred_sheet)
    if path.suffix.casefold() in {".csv", ".tsv"}:
        return _rows_from_csv(path)
    raise ValueError(f"不支持的输入格式：{path.suffix}")


def load_targets(path: str | Path) -> list[TargetProduct]:
    rows = _read_rows(Path(path), "ASIN输入")
    if not rows:
        raise ValueError("未读取到 ASIN 数据")
    targets: list[TargetProduct] = []
    seen_asins: set[str] = set()
    for row in rows:
        value = {field: first_nonempty(row, aliases) for field, aliases in TARGET_ALIASES.items()}
        asin = _validate_asin(value["asin"])
        if not asin:
            continue
        if asin in seen_asins:
            raise ValueError(f"重复 ASIN：{asin}")
        seen_asins.add(asin)
        price = to_float(value["selling_price_usd"])
        cpc = to_float(value["cpc_usd"])
        if price is None or cpc is None:
            raise ValueError(f"{asin} 缺少售价或 CPC")
        if price <= 0:
            raise ValueError(f"{asin} 售价必须大于 0")
        if cpc < 0:
            raise ValueError(f"{asin} CPC 不能小于 0")
        keywords = split_keywords(value["search_keywords"])
        targets.append(
            TargetProduct(
                asin=asin,
                selling_price_usd=price,
                cpc_usd=cpc,
                amazon_title=clean_text(value["amazon_title"]),
                search_keywords=keywords,
                material=clean_text(value["material"]),
                function=clean_text(value["function"]),
                structure=clean_text(value["structure"]),
                sales_unit=clean_text(value["sales_unit"]),
                pack_quantity=to_float(value["pack_quantity"]),
                accessory_quantity=to_float(value["accessory_quantity"]),
                length_cm=to_float(value["length_cm"]),
                width_cm=to_float(value["width_cm"]),
                height_cm=to_float(value["height_cm"]),
                product_weight_kg=to_float(value["product_weight_kg"]),
                package_length_cm=to_float(value["package_length_cm"]),
                package_width_cm=to_float(value["package_width_cm"]),
                package_height_cm=to_float(value["package_height_cm"]),
                package_weight_kg=to_float(value["package_weight_kg"]),
                fba_fee_usd=to_float(value["fba_fee_usd"]),
                notes=clean_text(value["notes"]),
            )
        )
    if not targets:
        raise ValueError("未读取到有效 ASIN 数据")
    return targets


def load_candidates(path: str | Path) -> list[CandidateProduct]:
    rows = _read_rows(Path(path), "1688候选")
    candidates: list[CandidateProduct] = []
    for row in rows:
        value = {
            field: first_nonempty(row, aliases) for field, aliases in CANDIDATE_ALIASES.items()
        }
        asin = _validate_asin(value["asin"])
        if not asin:
            continue
        comparison_price = to_float(value["comparison_price_cny"])
        if comparison_price is not None and comparison_price <= 0:
            raise ValueError(f"{asin} 完整销售单元采购价必须大于 0")
        candidates.append(
            CandidateProduct(
                asin=asin,
                search_keyword=clean_text(value["search_keyword"]),
                offer_id=clean_text(value["offer_id"]),
                title=clean_text(value["title"]),
                product_url=clean_text(value["product_url"]),
                image_url=clean_text(value["image_url"]),
                supplier=clean_text(value["supplier"]),
                shop_url=clean_text(value["shop_url"]),
                province=clean_text(value["province"]),
                city=clean_text(value["city"]),
                page_price_text=clean_text(value["page_price_text"]),
                comparison_price_cny=comparison_price,
                min_order_quantity=to_float(value["min_order_quantity"]),
                sales_text=clean_text(value["sales_text"]),
                return_rate=clean_text(value["return_rate"]),
                material=clean_text(value["material"]),
                function=clean_text(value["function"]),
                structure=clean_text(value["structure"]),
                sales_unit=clean_text(value["sales_unit"]),
                pack_quantity=to_float(value["pack_quantity"]),
                accessory_quantity=to_float(value["accessory_quantity"]),
                length_cm=to_float(value["length_cm"]),
                width_cm=to_float(value["width_cm"]),
                height_cm=to_float(value["height_cm"]),
                detail_text=clean_text(value["detail_text"]),
                screenshot_path=clean_text(value["screenshot_path"]),
                source_timestamp=clean_text(value["source_timestamp"]),
                manual_valid=clean_text(value["manual_valid"]),
                manual_reason=clean_text(value["manual_reason"]),
            )
        )
    return candidates
